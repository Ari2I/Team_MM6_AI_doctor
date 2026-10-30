"""Локальный генератор вымышленных заявок для тестирования Excel-импорта."""

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

DATA_DIR = Path(__file__).resolve().parent
HEADERS = (
    "Торговое наименование",
    "Дозировка",
    "Количество",
    "Сумма",
    "Диагноз (из заключения)",
)
MESSY_HEADERS = (
    "  тОРГОВОЕ   НАИМЕНОВАНИЕ  ",
    "\nДОЗИРОВКА  ",
    "  количество\n",
    "  сУмМа  ",
    " диагноз  (из\nзаключения)  ",
)

CellValue = str | int | float | None
SampleRow = tuple[str, str, int, int | float, str | None]

# По согласованию с пользователем выбраны существующие торговые наименования:
# ТЗ с примерами отсутствует. Все значения заявки вымышленные.
SAMPLE_ROWS: tuple[SampleRow, ...] = (
    ("Нурофен", "200 мг", 2, 1200, "Учебный диагноз 01"),
    ("Но-шпа", "40 мг", 1, 850.50, "Учебный диагноз 01"),
    ("Панадол", "500 мг", 3, 2100.75, "Учебный диагноз 02"),
    ("Конкор", "5 мг", 2, 640, "Учебный диагноз 02"),
    ("Омез", "20 мг", 4, 3200.25, "Учебный диагноз 03"),
    ("Супрастин", "25 мг", 1, 475.50, "Учебный диагноз 03"),
)


def write_workbook(
    path: Path,
    headers: Sequence[str],
    rows: Sequence[Sequence[CellValue]],
) -> None:
    """Записать один лист: заголовки в первой строке, далее данные без индекса."""
    workbook = Workbook()
    try:
        worksheet = workbook.active
        assert worksheet is not None
        worksheet.title = "Заявка"
        worksheet.append(list(headers))
        for row in rows:
            worksheet.append(list(row))

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=index)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.column_dimensions[get_column_letter(index)].width = (
                38 if "наименование" in header.casefold() or "диагноз" in header.casefold() else 20
            )
            if "".join(header.split()).casefold() == "сумма":
                for row_index in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row_index, column=index).number_format = "0.00"

        worksheet.row_dimensions[1].height = 42
        workbook.save(path)
    finally:
        workbook.close()


def generate_samples(rows: Sequence[SampleRow] = SAMPLE_ROWS, output_dir: Path = DATA_DIR) -> None:
    """Создать шесть файлов; каждый некорректный пример содержит один вид дефекта."""
    if not 5 <= len(rows) <= 8:
        raise ValueError("Для тестовой заявки требуется от 5 до 8 строк.")
    output_dir.mkdir(parents=True, exist_ok=True)
    write_workbook(output_dir / "sample_ok.xlsx", HEADERS, rows)
    write_workbook(
        output_dir / "sample_missing_column.xlsx",
        [header for index, header in enumerate(HEADERS) if index != 3],
        [tuple(value for index, value in enumerate(row) if index != 3) for row in rows],
    )
    bad_rows: list[Sequence[CellValue]] = list(rows)
    bad_rows[0] = (*rows[0][:3], "не число", rows[0][4])
    write_workbook(output_dir / "sample_bad_types.xlsx", HEADERS, bad_rows)
    write_workbook(output_dir / "sample_empty.xlsx", HEADERS, [])
    write_workbook(output_dir / "sample_messy_headers.xlsx", MESSY_HEADERS, rows)
    (output_dir / "sample_corrupt.xlsx").write_text(
        "Это обычный текстовый файл, а не книга Excel.\n", encoding="utf-8"
    )


if __name__ == "__main__":
    generate_samples()
