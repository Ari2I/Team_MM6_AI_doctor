"""Создать дополнительные вымышленные примеры, не изменяя исходные шесть файлов."""

from collections.abc import Sequence
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook

DATA_DIR = Path(__file__).resolve().parent
HEADERS = (
    "Торговое наименование",
    "Дозировка",
    "Количество",
    "Сумма",
    "Диагноз (из заключения)",
)
CellValue = str | int | float | None


def write_sample(filename: str, rows: Sequence[Sequence[CellValue]]) -> None:
    """Записать строки на первый лист; второй лист содержит постороннюю таблицу."""
    workbook = Workbook()
    try:
        sheet = workbook.active
        assert sheet is not None
        for row in rows:
            sheet.append(list(row))
        workbook.create_sheet("Другой лист").append(["Не заголовок заявки"])
        workbook.save(DATA_DIR / filename)
    finally:
        workbook.close()


def write_sheetless_sample() -> None:
    """Убрать листы из структуры XLSX для проверки ошибки чтения такой книги."""
    buffer = BytesIO()
    workbook = Workbook()
    try:
        workbook.save(buffer)
    finally:
        workbook.close()
    with (
        ZipFile(buffer) as source,
        ZipFile(DATA_DIR / "sample_parser_no_sheets.xlsx", "w") as target,
    ):
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                root = ElementTree.fromstring(content)
                sheets = root.find(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets"
                )
                assert sheets is not None
                sheets.clear()
                content = ElementTree.tostring(root)
            target.writestr(entry, content)


def generate_parser_samples() -> None:
    """Создать только дополнительные файлы с префиксом sample_parser_."""
    write_sample(
        "sample_parser_layout.xlsx",
        [
            ["Сумма"],  # Одиночное совпадение не должно перекрыть настоящие заголовки.
            *([[]] * 8),
            [*HEADERS, "Примечание"],  # Заголовки в строке 10.
            ["  Учебный\u00a0 препарат\n А ", "  5\n мг ", "1\u00a0234,50", "₽ 2 345,60", "NA"],
            [],
            [" \n\u00a0 ", " ", None, None, "\t"],
            ["Учебный препарат Б", None, -2, "не число", None],
            [None, None, None, None, None, "Учебное примечание"],
            [0, None, 0, -1.25, None],
        ],
    )
    write_sample(
        "sample_parser_duplicates.xlsx",
        [[*HEADERS, "Сумма"], ["Учебный препарат", None, 1, 5, None, 999]],
    )
    write_sample("sample_parser_no_headers.xlsx", [["Учебный документ"], ["Посторонний текст"]])
    write_sample("sample_parser_late_header.xlsx", [*([[]] * 10), HEADERS, ["Учебный препарат"]])
    write_sample("sample_parser_blank.xlsx", [])
    write_sheetless_sample()


if __name__ == "__main__":
    generate_parser_samples()
