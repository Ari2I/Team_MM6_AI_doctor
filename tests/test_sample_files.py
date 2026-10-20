"""Проверка готовых файлов для будущего парсера, без генерации во время тестов."""

from pathlib import Path
from zipfile import BadZipFile, is_zipfile

import pytest
from openpyxl import load_workbook

DATA_DIR = Path(__file__).resolve().parent / "data"
EXPECTED_HEADERS = (
    "Торговое наименование",
    "Дозировка",
    "Количество",
    "Сумма",
    "Диагноз (из заключения)",
)


def read_rows(filename: str) -> list[tuple[object, ...]]:
    """Прочитать файл с диска и закрыть его, не выводя содержимое в логи."""
    workbook = load_workbook(DATA_DIR / filename, read_only=True)
    try:
        assert workbook.sheetnames == ["Заявка"]
        return list(workbook["Заявка"].iter_rows(values_only=True))
    finally:
        workbook.close()


def test_sample_ok() -> None:
    rows = read_rows("sample_ok.xlsx")
    assert rows[0] == EXPECTED_HEADERS
    assert 5 <= len(rows) - 1 <= 8
    for name, dosage, quantity, amount, diagnosis in rows[1:]:
        assert isinstance(name, str) and name.strip()
        assert isinstance(dosage, str) and dosage.strip()
        assert type(quantity) is int and quantity > 0
        assert type(amount) in (int, float) and amount > 0
        assert diagnosis is None or isinstance(diagnosis, str)


def test_sample_missing_column() -> None:
    rows = read_rows("sample_missing_column.xlsx")
    expected = [
        tuple(value for index, value in enumerate(row) if index != 3)
        for row in read_rows("sample_ok.xlsx")
    ]
    assert rows == expected
    assert "Сумма" not in rows[0]


def test_sample_bad_types() -> None:
    rows = read_rows("sample_bad_types.xlsx")
    correct = read_rows("sample_ok.xlsx")
    assert rows[0] == correct[0]
    assert len(rows) == len(correct)
    assert isinstance(rows[1][3], str)
    assert rows[1][:3] == correct[1][:3]
    assert rows[1][4:] == correct[1][4:]
    assert rows[2:] == correct[2:]


def test_sample_empty() -> None:
    assert read_rows("sample_empty.xlsx") == [EXPECTED_HEADERS]


def test_sample_messy_headers() -> None:
    rows = read_rows("sample_messy_headers.xlsx")
    correct = read_rows("sample_ok.xlsx")
    assert rows[1:] == correct[1:]
    assert rows[0] != EXPECTED_HEADERS
    assert all(isinstance(header, str) for header in rows[0])
    assert tuple(" ".join(str(header).split()).casefold() for header in rows[0]) == tuple(
        header.casefold() for header in EXPECTED_HEADERS
    )
    assert any("\n" in str(header) for header in rows[0])


def test_sample_corrupt() -> None:
    path = DATA_DIR / "sample_corrupt.xlsx"
    assert path.read_text(encoding="utf-8").strip()
    assert not is_zipfile(path)
    with path.open("rb") as stream, pytest.raises(BadZipFile):
        load_workbook(stream)
